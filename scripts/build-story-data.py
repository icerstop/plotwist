"""Rebuild the chronological story library from downloaded, immutable source files.
Python dependencies: beautifulsoup4, lxml, openpyxl. No network or model API calls.
"""
from pathlib import Path
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
import csv, json, re, math, hashlib, zipfile, io, calendar
import openpyxl
from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning
import warnings
from build_technology_costs import add_cost_topics
from build_world_bank import add_world_bank_topics
warnings.filterwarnings('ignore',category=XMLParsedAsHTMLWarning)

ROOT=Path(__file__).resolve().parents[1]
RAW=ROOT/'research/stories/raw'
OUT=ROOT/'public/stories'
OUT.mkdir(parents=True,exist_ok=True)
SERIES={}; SOURCES={}; TOPICS={}; EVIDENCE={}
NOW=datetime.now(timezone.utc).isoformat()
def read(name): return json.loads((RAW/name).read_text(encoding='utf-8'))
def slug(s):
 import unicodedata
 return re.sub(r'[^a-z0-9]+','-',unicodedata.normalize('NFKD',s.replace('ł','l')).encode('ascii','ignore').decode().lower()).strip('-')
def source(name,label=None,url=None,note=''):
 sid=slug(name)
 if sid not in SOURCES:
  p=RAW/(name+'.receipt.json')
  receipt=json.loads(p.read_text(encoding='utf8')) if p.exists() else {'url':url,'retrievedAt':NOW,'extraction':'verified manual transcription from linked primary source'}
  SOURCES[sid]={**receipt,'id':sid,'name':label or name,'note':note,'license':'Original provider terms; source attribution required. See linked source.'}
 return sid
def topic(id,title,note='',status='ready'):
 TOPICS[id]={'id':id,'title':title,'note':note,'status':status,'seriesIds':[],'defaults':[]}
def series(topicId,key,name,unit,unitKey,frequency='annual',notes='',kind='observed',entity='',metric='',interpolation='linear'):
 sid=topicId+'--'+slug(key)
 if sid not in SERIES:
  SERIES[sid]={'id':sid,'topicId':topicId,'name':name,'entity':entity or name,'metric':metric or name,'unit':unit,'unitKey':unitKey,'frequency':frequency,'notes':notes,'kind':kind,'interpolation':interpolation,'points':[]}
  TOPICS[topicId]['seriesIds'].append(sid)
 return SERIES[sid]
def point(s,period,value,sid,precision='year',**extra):
 if re.fullmatch(r'\d{4}',str(period)):
  dt=f'{period}-12-31'; start=f'{period}-01-01'
 elif re.fullmatch(r'\d{4}-Q[1-4]',period):
  y,q=int(period[:4]),int(period[-1]);m=q*3;dt=f'{y}-{m:02d}-{calendar.monthrange(y,m)[1]}';start=f'{y}-{m-2:02d}-01'
 elif re.fullmatch(r'\d{4}-\d{2}',period):
  y,m=map(int,period.split('-'));dt=f'{period}-{calendar.monthrange(y,m)[1]}';start=f'{period}-01'
 else: dt=str(period);start=extra.pop('periodStart',None)
 s['points'].append({'date':dt,'period':str(period),'datePrecision':precision,'periodStart':start,'value':value,'sourceId':sid,**extra})
def numeric(v):
 try: return float(v) if v is not None and str(v).strip()!='' else None
 except (TypeError,ValueError): return None
def owid(topicId,file,unit=None,unitKey=None,notes='',entities=None,columns=None,frequency='annual'):
 meta=read(file+'.metadata.json'); sid=source(file+'.csv',meta['chart'].get('citation','OWID')+' · '+meta['chart']['title'],note='OWID processing; full original metadata retained in download package.')
 SOURCES[sid]['metadata']=meta
 rows=list(csv.DictReader((RAW/(file+'.csv')).open(encoding='utf-8-sig',newline='')))
 valueCols=[c for c in rows[0] if c not in ['Entity','Code','Year','Quarter','Day','World region according to OWID','Launch vehicle class']]
 for col in valueCols:
  if columns and col not in columns:continue
  cm=next((m for k,m in meta['columns'].items() if k==col or m.get('titleShort')==col),next(iter(meta['columns'].values())))
  for entity in sorted(set(r['Entity'] for r in rows)):
   if entities and entity not in entities:continue
   rr=[r for r in rows if r['Entity']==entity]
   if not any(numeric(r[col]) is not None for r in rr):continue
   s=series(topicId,file+'-'+entity+'-'+col,(entity+' · ' if entity!='World' else '')+col,unit or cm.get('unit',''),unitKey or cm.get('unit',''),frequency,notes or meta['chart'].get('note',''),entity=entity,metric=col)
   for r in rr:point(s,r.get('Year') or r.get('Quarter') or r.get('Day'),numeric(r[col]),sid,'year' if 'Year'in r else 'quarter')
 return rows

add_world_bank_topics(ROOT, topic, series, point, source, TOPICS, SOURCES)

topic('coffee','Kawa czy inwestycja?','Model matematyczny w dotychczasowym studiu. Nie istnieje źródłowy historyczny dataset dla hipotetycznej stałej stopy zwrotu.','simulation')

topic('battery','Baterie coraz tańsze','Pobrano historię ogniw, nie całych pakietów. Brak porównywalnej historii pakietów z tego źródła.','partial')
owid('battery','price-of-lithium-ion-battery-cells',unit='USD 2024 / kWh',unitKey='usd2024-kwh-cell',notes='Reprezentatywne ceny ogniw litowo-jonowych, ceny stałe 2024. Nie cena kompletnego pakietu baterii.')
topic('solar','Słońce i inne źródła odnawialne','LCOE nowych elektrowni, ceny stałe 2025. Ten zbiór nie obejmuje paliw kopalnych ani cen energii dla gospodarstw domowych.','partial')
owid('solar','levelized-cost-of-energy',unit='USD 2025 / kWh',unitKey='usd2025-kwh-lcoe',notes='LCOE nowych elektrowni. Ceny stałe 2025. Nie uwzględnia pełnego kosztu integracji z siecią.')
topic('storage','Terabajt za grosze','Najniższa historyczna cena odnotowana do danego roku; nie średnia cena rynkowa. Serie HDD, SSD, RAM i flash są rozdzielone.')
owid('storage','historical-cost-of-computer-memory-and-storage',unit='USD 2020 / TB',unitKey='usd2020-tb',notes='Najniższa cena historyczna według OWID/McCallum; ceny stałe 2020. TB, nie GB. Luki pozostają puste.')
topic('chips','Moore w liczbach','Zestaw OWID/Rupp pokazuje historyczny trend mikroprocesorów; nie identyfikuje pojedynczych chipów ani GPU. Ułamkowe liczby pozostawiono tak, jak publikuje źródło. Linia łączy nieregularne obserwacje; odcinki pomiędzy nimi są interpolacją wizualną, nie dodatkowymi pomiarami.','partial')
# Year precision is not annual frequency: this source supplies 39 selected years,
# so the chart should connect those observations without inserting empty years.
owid('chips','transistors-per-microprocessor',unit='tranzystory',unitKey='transistors',frequency='irregular',notes='Opracowany trend OWID/Rupp, nie katalog modeli CPU/GPU. Nieregularne obserwacje z dokładnością do roku, połączone linią. Wartości nie są zaokrąglane do fikcyjnych liczb fizycznych tranzystorów.')
topic('ev','Elektryczna zmiana warty','Dostępny udział BEV + PHEV łącznie. To nie jest udział samych BEV; brak prognoz. Metadane źródła mają niespójne oznaczenie edycji 2025/2026, zachowane w pakiecie.','partial')
owid('ev','electric-car-sales-share',unit='% nowych aut',unitKey='ev-bev-phev-sales-share',notes='BEV + PHEV. Nowa sprzedaż, nie udział w całej flocie. Źródło: IEA przez OWID.')
topic('work','Mniej godzin, więcej wartości?','Produktywność PPP i godziny pochodzą z Penn World Table 11.0. Godziny przed 1950 r. z Huberman/Minns dotyczą innej populacji, dlatego zachowano je jako osobne serie.')
owid('work','labor-productivity-per-hour-pennworldtable',unit='int. USD 2021 / h',unitKey='ppp2021-hour')
owid('work','annual-working-hours-per-worker',unit='h / pracownika / rok',unitKey='hours-worker-year')
for s in list(SERIES.values()):
 if s['topicId']=='work' and s['metric']=='Working hours per worker':
  historical=[p for p in s['points'] if p['date']<'1950'];s['points']=[p for p in s['points'] if p['date']>='1950'];s['notes']='Penn World Table: pracownicy i samozatrudnieni w całej gospodarce, od 1950 r.'
  if historical:
   h=series('work',s['id']+'-historical',s['entity']+' · godziny (1870–1938)',s['unit'],'hours-production-worker-year',notes='Huberman i Minns: pełnoetatowi pracownicy produkcyjni poza rolnictwem. Inna populacja niż PWT.',entity=s['entity'],metric='Historical production-worker hours');h['points']=historical

topic('nvidia','NVIDIA przed i po AI','Daty końca kwartałów fiskalnych. Druga seria obejmuje wszystkie obszary poza Data Center, a nie samo Gaming. W FY2027 NVIDIA połączyła pozostałe segmenty jako Edge Computing; zmianę raportowania opisano w metadanych.','partial')
meta=read('nvidia-source.json');nv=read('nvidia-values.json');sid=source('nvidia-values.json','NVIDIA / OWID · przychody kwartalne');SOURCES[sid]['metadata']=meta
zero=date.fromisoformat(meta['display']['zeroDay']);entities={r['id']:r['name'] for r in meta['dimensions']['entities']['values']}
for v,days,e in zip(nv['values'],nv['years'],nv['entities']):
 name=entities[e];s=series('nvidia',name,name,'mln USD / kwartał','usd-million-quarter','quarterly',notes=TOPICS['nvidia']['note'],entity='NVIDIA',metric=name)
 dt=zero+timedelta(days=days);point(s,dt.isoformat(),v/1e6,sid,'day',dateBasis='Reported fiscal-quarter end',fiscalYear=dt.year+1 if dt.month>1 else dt.year,fiscalQuarter={1:4,4:1,7:2,10:3}.get(dt.month))

topic('space','Tańszy dostęp do kosmosu','Koszt USD 2021/kg na LEO przypisany do roku pierwszego udanego startu rakiety, nie historia cennika każdego roku. Dostępne rekordy minimalnego kosztu i pojedyncze rakiety.')
spaceRows=owid('space','cost-space-launches-low-earth-orbit',unit='USD 2021 / kg LEO',unitKey='usd2021-kg-leo',frequency='event',columns=['Launch cost per kilogram of payload'],notes='Szacunek kosztu dedykowanego startu / maksymalna ładowność LEO; data = rok pierwszego udanego startu, nie data wyceny.')
sid=slug('cost-space-launches-low-earth-orbit.csv')
for cls in ['All','Small','Medium','Heavy']:
 s=series('space','record-'+cls,'Najniższy koszt · '+cls,'USD 2021 / kg LEO','usd2021-kg-leo','event','Minimum narastające po rakietach w zbiorze CSIS. Nie jest aktualną ofertą operatora.','derived',interpolation='step');best=math.inf
 for y in sorted(set(r['Year'] for r in spaceRows)):
  rr=[r for r in spaceRows if r['Year']==y and (cls=='All' or r['Launch vehicle class']==cls)]
  if not rr:continue
  winner=min(rr,key=lambda r:float(r['Launch cost per kilogram of payload']));v=float(winner['Launch cost per kilogram of payload'])
  if v<best:best=v;point(s,y,v,sid,vehicle=winner['Entity'])

# Primary financial statements: explicit table selections, no fuzzy value extraction.
def tables(company,year):
 name=f'{company}-{year}.html';soup=BeautifulSoup((RAW/name).read_text(encoding='utf8'),'lxml');return name,soup.find_all('table')
def tableRows(t): return [' | '.join(td.get_text(' ',strip=True) for td in tr.find_all(['td','th'],recursive=False) if td.get_text(' ',strip=True)) for tr in t.find_all('tr')]
def nums(t,minimum=0):
 # Ignore footnote numbers; financial figures are comma-grouped or decimals.
 clean=re.sub(r'\(\d\)','',t)
 return [float(x.replace(',','')) for x in re.findall(r'(?<![\w.])(?:\d{1,3}(?:,\d{3})+|\d+\.\d+|\d+)(?![\w.])',clean) if float(x.replace(',',''))>=minimum]

def appleYearEnd(y):
 dt=date(y,9,30)
 return dt-timedelta(days=(dt.weekday()-5)%7)
topic('apple','Akcje Apple vs iPhone','Roczne przychody iPhone 2008–2025 oraz sprzedaż sztuk 2008–2018. Apple zakończyło raportowanie sztuk; nie uzupełniono ich szacunkami. Przychody i ceny akcji mają inne jednostki. Wczesna kategoria obejmuje też produkty i usługi powiązane.')
appleRevenue=series('apple','iphone-revenue','iPhone · przychody','mln USD / rok','usd-million-year',notes='Rok fiskalny Apple kończący się we wrześniu. 2008–2010: iPhone and related products and services. Dane porównawcze przekształcone przez Apple w późniejszych raportach.',entity='Apple',metric='iPhone net sales')
appleUnits=series('apple','iphone-units','iPhone · sprzedane sztuki','mln sztuk / rok','million-units-year',notes='Oficjalna sprzedaż sztuk, zakończona po FY2018. Brak estymacji za kolejne lata.',entity='Apple',metric='iPhone units sold')
for year,idx in [(2010,90),(2013,89),(2016,97),(2019,49),(2022,28),(2025,27)]:
 name,tt=tables('apple',year);sid=source(name,f'Apple Form 10-K FY{year}');rr=tableRows(tt[idx]);EVIDENCE[name]=rr
 r=next(r for r in rr if r.startswith('iPhone'));values=nums(r,1000)[:3];assert len(values)==3,(name,r)
 for y,v in zip(range(year,year-3,-1),values):point(appleRevenue,appleYearEnd(y).isoformat(),v,sid,'fiscal-year',periodStart=(appleYearEnd(y-1)+timedelta(days=1)).isoformat(),fiscalYear=y,dateBasis='Fiscal year ends on the last Saturday of September')
 if year<=2016:
  idx2={2010:17,2013:23,2016:31}[year];rr=tableRows(tt[idx2]);EVIDENCE[name+'-units']=rr;rows=[r for r in rr if r.startswith('iPhone')];r=rows[-1];values=nums(r,1000)[:3]
  for y,v in zip(range(year,year-3,-1),values):point(appleUnits,appleYearEnd(y).isoformat(),v/1000,sid,'fiscal-year',fiscalYear=y)
name,tt=tables('apple',2018);sid=source(name,'Apple Form 10-K FY2018');rr=tableRows(tt[26]);EVIDENCE[name+'-units']=rr
values=nums([r for r in rr if r.startswith('iPhone')][-1],1000)[:2]
assert len(values)==2
for y,v in zip([2018,2017],values):point(appleUnits,appleYearEnd(y).isoformat(),v/1000,sid,'fiscal-year',fiscalYear=y)
stock=json.loads((ROOT/'public/market/AAPL.json').read_text(encoding='utf8'))
sid=source('apple-close','Yahoo Finance · Apple daily close',stock['sourceUrl'],note='Existing downloaded market snapshot. Close adjusted for stock splits, not dividends.')
SOURCES[sid].update(retrievedAt=stock['retrievedAt'],extraction='Existing downloaded AAPL market snapshot')
s=series('apple','stock-close','Apple · cena zamknięcia','USD / akcję','usd-share','daily',notes='Cena zamknięcia skorygowana o splity, bez reinwestycji dywidend. Dni bez sesji nie są dopisywane.',entity='Apple',metric='Split-adjusted daily close')
for r in stock['rows']:point(s,r[0],r[1],sid,'day')
annualClose=series('apple','stock-fiscal-close','Apple · kurs na koniec roku fiskalnego','USD / akcję','usd-share','annual',notes='Ostatnia sesja w roku fiskalnym Apple; data osi = koniec roku fiskalnego. Rzeczywista data sesji w tradingDate. Close po splitach, bez dywidend.',entity='Apple',metric='Fiscal-year-end split-adjusted close')
for p in appleRevenue['points']:
 r=next(r for r in reversed(stock['rows']) if r[0]<=p['date']);point(annualClose,p['date'],r[1],sid,'fiscal-year',fiscalYear=p['fiscalYear'],tradingDate=r[0])

topic('spotify','Streaming zmienił muzykę','Roczne obserwacje 2017–2025: MAU, Premium oraz miesięczne ARPU uśrednione za rok. Nie należy dzielić rocznych przychodów przez stan subskrybentów na koniec roku i nazywać wyniku ARPU.')
spotSpecs=[('mau','MAUs','Spotify · MAU','mln MAU','million-mau'),('premium','Premium Subscribers','Spotify · Premium','mln subskrybentów','million-subscribers'),('arpu','Premium ARPU','Spotify · Premium ARPU','EUR / miesiąc','eur-subscriber-month')]
for year,indices in [(2019,[126,127,129]),(2022,[12,13,15]),(2025,[12,13,15])]:
 name,tt=tables('spotify',year);sid=source(name,f'Spotify Form 20-F {year}')
 for spec,idx in zip(spotSpecs,indices):
  key,label,title,unit,uk=spec;rr=tableRows(tt[idx]);EVIDENCE[name+'-'+key]=rr;r=next(r for r in rr if r.startswith(label));vv=nums(r)[:3];assert len(vv)==3
  s=series('spotify',key,title,unit,uk,entity='Spotify',metric=label,notes='MAU/Premium: stan na 31 grudnia. ARPU: miesięczny przychód Premium na średniego subskrybenta, średnia za cały rok.')
  for y,v in zip(range(year,year-3,-1),vv):point(s,str(y),v,sid)
 # Consolidated income statement annual revenue; select the table with three reporting years.
 candidates=[t for t in tt if 'Revenue' in t.get_text() and 'Cost of revenue' in t.get_text() and ('Net income' in t.get_text() or 'Net loss' in t.get_text()) and len(t.get_text())<12000]
 t=next(t for t in candidates if all(str(y) in t.get_text() for y in [year,year-1,year-2]));rr=tableRows(t);r=next(r for r in rr if r.startswith('Revenue'));vv=nums(r,1000)[:3];assert len(vv)==3;EVIDENCE[name+'-revenue']=rr
 s=series('spotify','revenue','Spotify · przychody','mln EUR / rok','eur-million-year',entity='Spotify',metric='Revenue')
 for y,v in zip(range(year,year-3,-1),vv):point(s,str(y),v,sid)

topic('cloud','Wyścig chmurowych gigantów','AWS i Google Cloud: przychody roczne. Azure: tylko oficjalnie ujawnione progi >75 i >100 mld USD dla FY2025/26, jako oddzielna seria dolnych granic. Brak wymyślonej historii Azure.','partial')
for company,ixs in [('amazon',[(2016,162),(2019,170),(2022,69),(2025,69)]),('alphabet',[(2019,273),(2022,94),(2025,94)])]:
 s=series('cloud',company,'AWS' if company=='amazon' else 'Google Cloud','mln USD / rok','usd-million-year',entity=company,metric='Annual segment revenue',notes='Rok kalendarzowy; różny zakres usług segmentów. Google Cloud obejmuje też Workspace.')
 for year,idx in ixs:
  name,tt=tables(company,year);sid=source(name,f'{company.title()} Form 10-K {year}');rr=tableRows(tt[idx]);EVIDENCE[name]=rr
  if company=='amazon':
   start=next(i for i,r in enumerate(rr) if r=='AWS');r=next(r for r in rr[start+1:] if r.startswith('Net sales'))
  else:r=next(r for r in rr if r.startswith('Google Cloud'))
  vv=nums(r,1000)[:3];assert len(vv)==3,(name,r)
  for y,v in zip(range(year-2,year+1),vv):point(s,str(y),v,sid)
s=series('cloud','azure-lower-bound','Azure · dolna granica przychodów','mln USD / rok fiskalny','usd-million-fiscal-year','annual',notes='Komunikaty podają „przekroczyły”, a nie dokładną wartość. Rok kończy się 30 czerwca.',kind='lower-bound',entity='Microsoft',metric='Azure revenue lower bound')
for y,v in [(2025,75000),(2026,100000)]:point(s,f'{y}-06-30',v,source(f'microsoft-{y}.html',f'Microsoft FY{y} Q4'),'fiscal-year',fiscalYear=y,valueQualifier='>',periodStart=f'{y-1}-07-01')

topic('ai-cost','Ile kosztuje odpowiedź AI?','Historyczny minimalny koszt osiągnięcia danego progu benchmarku, według Epoch AI. Cena mieszana 3:1 wejście:wyjście; nie koszt całego zadania ani aktualny cennik. Różne progi benchmarków pozostają osobnymi seriami.')
sid=source('epoch-price-frontier.csv','Epoch AI / Artificial Analysis · historical inference frontier')
for r in csv.DictReader((RAW/'epoch-price-frontier.csv').open(encoding='utf8')):
 key=r['Benchmark']+'-'+r['Threshold model'];s=series('ai-cost',key,r['Benchmark']+' ≥ '+r['Threshold model'],'USD / mln tokenów (3:1)','usd-million-tokens-3to1','event',notes='Próg: '+r['Performance range']+'. Średnia ważona cen input/output 3:1. Model i wynik zachowane przy każdej obserwacji.',kind='derived',interpolation='step')
 point(s,r['Release Date'],float(r['USD per 1M Tokens']),sid,'day',model=r['Model Name'],benchmark=r['Benchmark'],threshold=r['Performance range'],score=numeric(r.get('Benchmark score')))

topic('patents','Mapa wynalazków','Wnioski patentowe rezydentów (WIPO przez World Bank), liczby bezwzględne i na milion mieszkańców. To nie liczba udzielonych patentów. Można łączyć z już dostępnymi nakładami B+R.')
pat,pop=read('wb-patents.json'),read('wb-population.json');assert len(pat[1])==pat[0]['total'] and len(pop[1])==pop[0]['total']
population={(r['countryiso3code'],r['date']):r['value'] for r in pop[1]};sid=source('wb-patents.json','WIPO / World Bank · resident patent applications');popid=source('wb-population.json','World Bank · population')
for r in pat[1]:
 if not r['countryiso3code']:continue
 c=r['countryiso3code'];s=series('patents',c+'-count',r['country']['value']+' · zgłoszenia','zgłoszenia / rok','patent-applications-year',entity=c,metric='Resident patent applications');point(s,r['date'],r['value'],sid)
 p=population.get((c,r['date']));s=series('patents',c+'-per-million',r['country']['value']+' · na mln osób','zgłoszenia / mln osób','patent-applications-million-pop','annual','Zgłoszenia rezydentów / ludność w tym samym roku × 1 000 000.','derived',entity=c,metric='Resident patent applications per million')
 point(s,r['date'],r['value']/p*1e6 if r['value'] is not None and p else None,sid,additionalSourceIds=[popid])

topic('inflation','Co kupi dzisiejsze 100 zł?','Roczne CPI GUS od 1950 r. Łańcuch cen obliczono mnożąc indeksy rok do roku; nie sumując inflacji. Kwota 100 zł oznacza siłę nabywczą współczesnej jednostki pieniężnej w cenach roku bazowego, nie historyczny banknot sprzed denominacji.')
sid=source('gus-cpi.csv','GUS · roczny CPI od 1950 r.');rows=list(csv.DictReader((RAW/'gus-cpi.csv').open(encoding='cp1250'),delimiter=';'))
chain=100
for r in rows:
 year=int(r['Rok']);v=float(r['Wartość'].replace(',','.'));s=series('inflation','cpi-yoy','Polska · CPI rok poprzedni = 100','indeks r/r','cpi-yoy',notes='Średnioroczny poziom cen w relacji do poprzedniego roku.');point(s,str(year),v,sid)
 if year==1950:chain=100
 else:chain*=v/100
 s=series('inflation','cpi-chain','Polska · poziom cen (1950 = 100)','indeks 1950 = 100','cpi-chain-1950','annual','Iloczyn rocznych indeksów GUS; średnioroczne ceny.','derived');point(s,str(year),chain,sid)
# A modern base avoids silently interpreting pre-denomination złoty as today's currency.
base=next(p['value'] for p in SERIES['inflation--cpi-chain']['points'] if p['period']=='2000')
for p in SERIES['inflation--cpi-chain']['points']:
 if p['period']<'2000':continue
 s=series('inflation','purchasing-power','Siła nabywcza 100 zł (ceny 2000)','zł w cenach 2000','pln2000-purchasing-power','annual','100 × CPI(2000) / CPI(t). Roczne średnie; rok 2000 = 100 zł.','derived');point(s,p['period'],100*base/p['value'],sid)

topic('wages','Pensja vs metr mieszkania','17 miast. Ceny transakcyjne NBP, rynek pierwotny i wtórny. Płace GUS BDL: brutto, podmioty 10+ i sfera budżetowa. Wskaźnik dostępności: średnia z 4 kwartałów cen / średnia miesięczna płaca w tym samym mieście i roku. Nie jest czasem oszczędzania po kosztach życia.')
housing=openpyxl.load_workbook(RAW/'nbp-housing.xlsx',data_only=True);sid=source('nbp-housing.xlsx','NBP · BaRN ceny transakcyjne');housingAnnual=defaultdict(list)
for sheet,market in [(housing.worksheets[1],'pierwotny'),(housing.worksheets[2],'wtórny')]:
 header=[c.value for c in sheet[7]];cities=[str(c).replace('*','') for c in header[24:41]]
 for row in sheet.iter_rows(min_row=8,max_col=41,values_only=True):
  label=row[23];m=re.fullmatch(r'(I|II|III|IV) (\d{4})',str(label or ''))
  if not m:continue
  q=['I','II','III','IV'].index(m[1])+1;y=int(m[2])
  for city,v in zip(cities,row[24:41]):
   v=numeric(v);s=series('wages',city+'-'+market+'-price',city+' · '+market+' · cena m²','PLN / m²','pln-m2','quarterly',entity=city,metric='Cena transakcyjna m² · '+market);point(s,f'{y}-Q{q}',v,sid,'quarter')
   if v is not None:housingAnnual[(city,market,y)].append(v)
wages={};unitIds=set();bdlRows=[]
for file in ['bdl-wages.json','bdl-wages-1.json','bdl-wages-2.json','bdl-wages-3.json']:
 d=read(file);bdlRows.extend(d['results']);wageSource=source(file,'GUS BDL · wynagrodzenia brutto · 64428')
 for r in d['results']:
  assert r['id'] not in unitIds;unitIds.add(r['id'])
  city=next((c for c in cities if r['name'] in ['Powiat m.'+c,'Powiat m. '+c,'Powiat m.st. '+c,'Powiat m. st. '+c]),None)
  if not city:continue
  s=series('wages',city+'-wage',city+' · płaca brutto','PLN / miesiąc','pln-month','annual','GUS BDL 64428: podmioty 10+ pracujących i jednostki budżetowe. Nie cała populacja zatrudnionych.',entity=city,metric='Średnie miesięczne wynagrodzenie brutto')
  for p in r['values']:
   wages[(city,int(p['year']))]=(p['val'],wageSource);point(s,str(p['year']),p['val'],wageSource,attributeId=p.get('attrId'),territoryId=r['id'])
assert len(bdlRows)==381,(len(bdlRows),'Incomplete BDL pagination')
assert len(set(c for c,y in wages))==17,sorted(set(c for c,y in wages))
for (city,market,y),vv in housingAnnual.items():
 if len(vv)!=4 or (city,y) not in wages:continue
 pay,wageSource=wages[(city,y)];avg=sum(vv)/4
 s=series('wages',city+'-'+market+'-affordability',city+' · '+market+' · pensje za m²','mies. płacy brutto / m²','gross-monthly-salaries-m2','annual','Prosta średnia 4 kwartalnych cen transakcyjnych / średnia miesięczna płaca brutto tego samego miasta i roku.','derived',entity=city,metric='Cena m² / miesięczna płaca brutto')
 point(s,str(y),avg/pay if pay else None,sid,additionalSourceIds=[wageSource],annualPrice=avg,monthlyGrossWage=pay)

topic('gaming','Konsole na przestrzeni lat','Oficjalne roczne dostawy sprzętu Nintendo w latach fiskalnych kończących się 31 marca. Nie sell-through. Starsze konsole nie mają kompletnej historii od premiery; skumulowane serie tworzone tylko od pierwszego roku sprzedaży. Kolumny zbiorcze „2015~” i „2024~” pominięto, bo nie oznaczają jednego roku. Brak kompletnej historii Sony/Microsoft.','partial')
book=openpyxl.load_workbook(RAW/'nintendo-sales.xlsx',data_only=True);sid=source('nintendo-sales.xlsx','Nintendo · Consolidated Sales Transition · March 2026')
for sheet in book.worksheets:
 rows=list(sheet.values);head=rows[3];yearCols=[(i,int(str(v)[-4:])) for i,v in enumerate(head) if re.fullmatch(r'FY3/\d{4}',str(v))];model=None;hw=False;release=None
 for row in rows[5:]:
  if row[0] and not str(row[0]).lstrip().startswith('[Note]'):model=str(row[0]).strip();hw=False;release=None
  if row[1] and 'released on' in str(row[1]):release=str(row[1])
  first=yearCols[0][0];cells=[str(v).strip() for v in row[:first] if v is not None]
  if 'Hardware' in cells:hw=True
  if any(v.startswith('of which') or v=='Software' for v in cells):hw=False
  if not hw or 'Total' not in cells or not model:continue
  s=series('gaming',model+'-annual',model+' · dostawy roczne','mln sztuk / rok','million-units-year','annual','Zaokrąglone przez Nintendo do 10 tys. Komórki ±0,1 mają format wyświetlania 0/-0 i oznaczają wielkości poniżej dokładności tabeli; pokazujemy je jako ≈0, zachowując surową wartość i znak w metadanych. '+(release or ''),entity=model,metric='Annual hardware shipments')
  for col,y in yearCols:
   v=numeric(row[col]);point(s,f'{y}-03-31',round(v)/100 if v is not None else None,sid,'fiscal-year',periodStart=f'{y-1}-04-01',fiscalYear=y,rawSpreadsheetValue=v,valueQualifier='approximately',roundingUnits=10000)
  # Only newer systems are fully covered from their launch year.
  if model in ['Nintendo DS','Wii','Nintendo 3DS','Wii U','Nintendo Switch','Nintendo Switch 2']:
   c=series('gaming',model+'-cumulative',model+' · dostawy narastająco','mln sztuk','million-units-stock','annual','Suma zaokrąglonych rocznych dostaw; może różnić się od oficjalnej wartości life-to-date wskutek zaokrągleń.','derived',entity=model,metric='Cumulative rounded shipments');total=0
   for p in s['points']:
    if p['value'] is not None:total+=p['value'];point(c,p['date'],round(total,6),sid,'fiscal-year',fiscalYear=p['fiscalYear'],valueQualifier='approximately')

topic('adoption','Adopcja aplikacji · udokumentowane punkty','Oddzielono MAU, WAU i regiony. Facebook: roczne MAU 2011–2023, później firma zakończyła raportowanie tej miary. Instagram, ChatGPT i TikTok: komunikaty o kamieniach milowych, nie ciągłe pomiary. Brak wspólnej definicji pozwalającej na uczciwy wyścig wszystkich aplikacji do 100 mln.','partial')
sid=source('meta-2012.html','Facebook Form 10-K 2012');s=series('adoption','facebook-mau','Facebook · MAU','mln MAU','million-mau','annual','Stan na 31 grudnia. Wartości zaokrąglone przez firmę; definicja obejmowała także zarejestrowanych użytkowników Facebooka korzystających z Messengera. Zmiany definicji i ograniczenia pomiaru w raportach.',entity='Facebook',metric='Monthly active users',interpolation='step')
for y,v in [(2011,845),(2012,1060)]:point(s,str(y),v,sid)
for y in range(2013,2024):
 name=f'meta-{y}.html';text=BeautifulSoup((RAW/name).read_text(encoding='utf8'),'lxml').get_text(' ',strip=True)
 match=re.search(r'As of December\s+31,\s*'+str(y)+r'\s*,\s*we had\s+([\d.]+)\s+billion MAUs',text,re.I)
 assert match,(name,'Missing annual MAU statement')
 sid=source(name,f'Meta / Facebook Form 10-K {y}');EVIDENCE[name+'-mau']=[match[0]]
 point(s,str(y),round(float(match[1])*1000,3),sid,valueQualifier='approximately')
s=series('adoption','instagram-mau','Instagram · MAU','mln MAU','million-mau','event','Globalne kamienie milowe z oficjalnych komunikatów Meta. Daty komunikatów, nie ustalony dzień przekroczenia progu. Brak danych między komunikatami.',entity='Instagram',metric='Monthly active users',interpolation='step')
for dt,v,file,prec in [('2015-09',400,'instagram-400.html','month'),('2016-06',500,'instagram-500.html','month'),('2016-12',600,'instagram-700.html','month'),('2017-04-26',700,'instagram-700.html','day'),('2017-09-26',800,'instagram-800.html','day'),('2018-06-20',1000,'meta-timeline.html','day')]:
 point(s,dt,v,source(file,'Meta Newsroom · Instagram MAU'),'month' if prec=='month' else 'day',valueQualifier='approximately',dateBasis='Milestone announcement, not a daily measurement')
sid=source('openai-adoption.pdf','OpenAI · ChatGPT usage and adoption patterns at work, p.6');s=series('adoption','chatgpt-wau','ChatGPT · WAU','mln WAU','million-wau','event','Dwa zaokrąglone kamienie milowe z raportu, miesiąc pomiaru z osi wykresu. Nie MAU; nie estymujemy wartości między punktami.',entity='ChatGPT',metric='Weekly active users',interpolation='step')
point(s,'2023-11',100,sid,'month',valueQualifier='approximately');point(s,'2025-07',700,sid,'month',valueQualifier='>')
sid=source('openai-axios','OpenAI · Axios partnership','https://openai.com/index/partnering-with-axios-expands-openai-work-with-the-news-industry/');point(s,'2025-01-15',300,sid,'day',valueQualifier='>',dateBasis='Publication date of milestone')
sid=source('tiktok-billion.html','TikTok · Thanks a billion');s=series('adoption','tiktok-mau','TikTok · MAU (globalnie)','mln MAU','million-mau','event','Pojedynczy globalny kamień milowy, nie seria czasowa.',entity='TikTok',metric='Monthly active users',interpolation='step');point(s,'2021-09-27',1000,sid,'day',valueQualifier='>')

topic('internet-speed','Ile trwa pobranie filmu?','Archiwalne pomiary Ofcom dla Wielkiej Brytanii: mediana realnej przepustowości domowego łącza do routera, 2019 i 2023. To nie historia reklamowanych prędkości ani modemów od lat 90.','partial')
sid=source('ofcom-observations','Ofcom · home broadband performance','https://www.ofcom.org.uk/phones-and-broadband/bills-and-charges/ban-on-inflation-linked-mid-contract-price-rise',note='Verified transcription: November 2019 42.1 Mbit/s and March 2023 69.4 Mbit/s. Methodology: March 2023 home broadband performance report.')
s=series('internet-speed','download','UK · mediana pobierania','Mbit/s','mbit-second','event',notes='Mediana średniej dobowej realnej prędkości pobierania. Dwa udokumentowane pomiary; nie symulacja postępu w latach pośrednich.',entity='United Kingdom')
for dt,v in [('2019-11',42.1),('2023-03',69.4)]:point(s,dt,v,sid,'month')
s=series('internet-speed','download-time','Pobranie 1 GB · idealny czas','sekundy','seconds','event','8 000 megabitów / Mbit/s. 1 GB dziesiętny; bez narzutu i ograniczeń serwera.','derived',entity='United Kingdom')
for dt,v in [('2019-11',42.1),('2023-03',69.4)]:point(s,dt,8000/v,sid,'month')

add_cost_topics(RAW, topic, series, point, source, TOPICS, SOURCES)

# Prune entirely empty series, sort chronologically, enforce unique observation dates.
for sid in list(SERIES):
 s=SERIES[sid]
 if not any(p['value'] is not None for p in s['points']):
  TOPICS[s['topicId']]['seriesIds'].remove(sid);del SERIES[sid];continue
 s['points'].sort(key=lambda p:p['date']);seen=set()
 for p in s['points']:
  date.fromisoformat(p['date']);assert p['date'] not in seen,(sid,p['date']);seen.add(p['date'])
  assert p['value'] is None or math.isfinite(p['value']),(sid,p)
  assert p['sourceId'] in SOURCES
  assert p['date']<=date.today().isoformat() or (p['value'] is None and p['datePrecision']=='year'),(sid,p)
 obs=[p for p in s['points'] if p['value'] is not None]
 s.update(count=len(obs),missing=len(s['points'])-len(obs),start=obs[0]['date'],end=obs[-1]['date'],sourceIds=sorted({sid for p in s['points'] for sid in [p['sourceId'],*p.get('additionalSourceIds',[])]}),datePrecisions=sorted(set(p['datePrecision'] for p in s['points'])))

defaultKeys={'apple':['iphone-revenue','stock-fiscal-close'],'nvidia':['data-centers-and-ai','gaming-devices-automotive'],'spotify':['mau'],'cloud':['amazon','alphabet'],'inflation':['purchasing-power'],'adoption':['facebook-mau','instagram-mau'],'internet-speed':['download'],'gaming':['nintendo-switch-cumulative','wii-cumulative'],'space':['record-all'],'wages':['warszawa-wtorny-affordability'],'patents':['pol-per-million','usa-per-million'],
 'solar':['levelized-cost-of-energy-world-solar-photovoltaic','levelized-cost-of-energy-world-onshore-wind'],
 'storage':['historical-cost-of-computer-memory-and-storage-world-disk','historical-cost-of-computer-memory-and-storage-world-solid-state'],
 'ev':['electric-car-sales-share-china-share-of-new-cars-that-are-electric','electric-car-sales-share-united-states-share-of-new-cars-that-are-electric','electric-car-sales-share-poland-share-of-new-cars-that-are-electric'],
 'work':['labor-productivity-per-hour-pennworldtable-poland-productivity-output-per-hour-worked','labor-productivity-per-hour-pennworldtable-germany-productivity-output-per-hour-worked']}
for t in TOPICS.values():
 if t['id'] in defaultKeys:t['defaults']=[t['id']+'--'+k for k in defaultKeys[t['id']] if t['id']+'--'+k in SERIES]
 if not t['defaults']:t['defaults']=[sid for sid in t['seriesIds'] if SERIES[sid]['count']>=2][:1]
 t['count']=sum(SERIES[s]['count'] for s in t['seriesIds']);t['seriesCount']=len(t['seriesIds'])
 t['start']=min((SERIES[s]['start'] for s in t['seriesIds']),default=None);t['end']=max((SERIES[s]['end'] for s in t['seriesIds']),default=None)
 t.setdefault('defaultMode','index' if t['id']=='apple' else 'native')
 t.setdefault('defaultChart','cards' if t['id']=='adoption' else 'line')
for src in SOURCES.values():
 n=src['name'].lower()
 labels=[('yahoo','Yahoo Finance'),('apple','Apple / SEC'),('spotify','Spotify / SEC'),('amazon','Amazon / SEC'),('alphabet','Alphabet / SEC'),('meta','Meta / SEC'),('facebook','Meta / SEC'),('nvidia','NVIDIA / OWID'),('nintendo','Nintendo'),('microsoft','Microsoft'),('gus','GUS'),('nbp','NBP'),('openai','OpenAI'),('tiktok','TikTok'),('ofcom','Ofcom'),('epoch','Epoch AI'),('wipo','WIPO / World Bank'),('world bank','World Bank'),('irena','IRENA / OWID'),('iea','IEA / OWID')]
 src['shortName']=src.get('shortName') or ('NHGRI' if 'nhgri' in n else next((label for word,label in labels if word in n),'OWID'))
angles=json.loads((ROOT/'research/stories/story-notes.json').read_text(encoding='utf8'))
for t in TOPICS.values():t['variants']=t.get('variants') or angles.get(t['id'],['Porównanie wybranych krajów w czasie: linie, karty lub wyścig słupków. Zakres i luki sprawdź dla każdej serii.'] if t['seriesCount'] else ['Model hipotetyczny w module Studio; historyczne symulacje wpłat w module Giełda.'])
manifest={'schemaVersion':1,'builtAt':NOW,'topics':list(TOPICS.values()),'series':[{k:v for k,v in s.items() if k!='points'} for s in SERIES.values()],'sources':list(SOURCES.values())}
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,separators=(',',':')),encoding='utf8')
(ROOT/'src/story-catalog.json').write_text(json.dumps({'builtAt':NOW,'topics':[{k:v for k,v in t.items() if k not in ['seriesIds','defaults']} for t in TOPICS.values()],'observations':sum(s['count'] for s in SERIES.values()),'seriesCount':len(SERIES)},ensure_ascii=False,separators=(',',':')),encoding='utf8')
headers=['topic_id','series_id','series_name','entity','metric','date','period','date_precision','period_start','value','unit','frequency','kind','source_url','retrieved_at','attributes_json']
def csvRows(ss):
 for s in ss:
  for p in s['points']:
   src=SOURCES[p['sourceId']];attrs={k:v for k,v in p.items() if k not in ['date','period','datePrecision','periodStart','value','sourceId']}
   yield [s['topicId'],s['id'],s['name'],s['entity'],s['metric'],p['date'],p['period'],p['datePrecision'],p.get('periodStart'),p['value'],s['unit'],s['frequency'],s['kind'],src['url'],src['retrievedAt'],json.dumps(attrs,ensure_ascii=False)]
def writeCsv(path,rows):
 with path.open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.writer(f);w.writerow(headers);w.writerows(rows)
for t in TOPICS.values():
 ss=[SERIES[s] for s in t['seriesIds']];(OUT/(t['id']+'.json')).write_text(json.dumps({'topic':t,'series':ss},ensure_ascii=False,separators=(',',':')),encoding='utf8')
 writeCsv(OUT/(t['id']+'.csv'),sorted(csvRows(ss),key=lambda r:(r[5],r[1])))
# The complete CSV is streamed into the ZIP to avoid oversized web assets.
(OUT/'all-observations.csv').unlink(missing_ok=True)
(ROOT/'research/stories/extracted-financial-tables.json').write_text(json.dumps(EVIDENCE,ensure_ascii=False,indent=2),encoding='utf8')
quality={'builtAt':NOW,'topics':len(TOPICS),'series':len(SERIES),'observations':sum(s['count'] for s in SERIES.values()),'missingCells':sum(s['missing'] for s in SERIES.values()),'duplicateDates':0,'nonFiniteValues':0,'coverage':[{'id':t['id'],'status':t['status'],'series':t['seriesCount'],'observations':t['count'],'start':t['start'],'end':t['end'],'limitations':t['note']} for t in TOPICS.values()]}
(OUT/'quality-report.json').write_text(json.dumps(quality,ensure_ascii=False,indent=2),encoding='utf8')
readme='''# Plotwist — biblioteka historii\n\nSnapshot źródeł, nie bieżący feed. manifest.json opisuje serie, źródła i ograniczenia każdego tematu.\nPełny all-observations.csv jest w archiwum ZIP; pojedyncze tematy mają osobne CSV i JSON.\nCSV ma format długi: jeden wiersz = seria × okres; pusta wartość oznacza brak danych, nigdy zero.\nDaty są rosnące. date_precision i period zachowują dokładność źródła. Roczne daty 31 grudnia i kotwice fiskalne służą porządkowaniu, nie oznaczają pomiaru dziennego.\nNie sumuj stanów, procentów i przepływów. Nie mieszaj USD o różnych latach cenowych, MAU/WAU ani BEV/BEV+PHEV.\nkind=derived oznacza obliczenie; lower-bound oznacza dolną granicę, nie dokładny wynik.\nPunkty przejściowe animacji nie są nowymi obserwacjami.\nLicencje oryginalnych dostawców zachowują moc; metadane OWID i bezpośrednie URL źródeł są w pakiecie.\nOdtwarzanie: scripts/build-story-data.py; pobieranie: scripts/fetch-story-sources.mjs, fetch-story-files.mjs, fetch-story-filings.mjs.\n'''
(OUT/'README.md').write_text(readme,encoding='utf8')
guide=readme+'\n## Tematy i dostępna historia\n\n'
for t in TOPICS.values():
 guide+=f"### {t['title']}\n\n{t['seriesCount']} serii; {t['count']} obserwacji; {t['start'] or '—'} → {t['end'] or '—'}. Status: {t['status']}.\n\n{t['note']}\n\n"+'\n'.join('- '+a for a in t['variants'])+'\n\n'
(OUT/'STORY-GUIDE.md').write_text(guide,encoding='utf8')
(OUT/'TECHNOLOGY-COSTS.md').write_text((ROOT/'research/stories/technology-costs.md').read_text(encoding='utf8'),encoding='utf8')
(OUT/'WORLD-BANK.md').write_text((ROOT/'research/world-bank/README.md').read_text(encoding='utf8'),encoding='utf8')
with zipfile.ZipFile(OUT/'plotwist-story-datasets.zip','w',compression=zipfile.ZIP_DEFLATED) as z:
 with z.open('all-observations.csv','w',force_zip64=True) as entry:
  with io.TextIOWrapper(entry,encoding='utf-8-sig',newline='') as text:
   writer=csv.writer(text);writer.writerow(headers);writer.writerows(sorted(csvRows(SERIES.values()),key=lambda r:(r[5],r[1])))
 for p in OUT.iterdir():
  if p.suffix in ['.json','.md']:z.write(p,p.name)
 for p in RAW.glob('*.metadata.json'):z.write(p,'source-metadata/'+p.name)
 for pattern in ['epoch-thought-*','epoch-chip-performance.csv*','nhgri-sequencing-costs.xls*']:
  for p in RAW.glob(pattern):z.write(p,'technology-cost-sources/'+p.name)
 for p in RAW.glob('gdp-per-capita-maddison-project-database.*'):z.write(p,'historical-economy-sources/'+p.name)
 z.write(ROOT/'research/world-bank/receipt.json','world-bank-sources/receipt.json')
 z.write(ROOT/'research/stories/extracted-financial-tables.json','evidence/extracted-financial-tables.json')
print(json.dumps({k:v for k,v in quality.items() if k!='coverage'},ensure_ascii=False))
for t in TOPICS.values():print(t['id'],t['seriesCount'],t['count'],t['start'],t['end'],t['status'])
