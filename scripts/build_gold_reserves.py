"""Audited Poland gold data. Missing months stay missing; no rounded ECB volumes."""
import calendar
import csv
import json
import hashlib
import re
from decimal import Decimal
from pathlib import Path
from bs4 import BeautifulSoup
import openpyxl

TONNES_PER_MILLION_OZT = Decimal('31.1034768')
TOPIC = 'poland-gold'
COVERAGE = 'Rezerwy w tonach: kompletna historia miesięczna 2000-01–2026-08. Brakujące wcześniej miesiące uzupełniono raportowanymi danymi NBP z publikacji PAP, prasy finansowej i raportu bankowego; pochodzenie oraz precyzja są zapisane przy pomiarze. Nie użyto interpolacji ani szacunków z ceny złota. Serie wartościowe nadal mają lukę EBC 2014-01–2014-08. Stan obejmuje złoto NBP także przechowywane za granicą, nie zasoby geologiczne Polski.'

def reported_volume(root):
    return json.loads((root/'research/gold-reserves/raw/reported-monthly-volume.json').read_text(encoding='utf8'))

def months(first, last):
    y, m = map(int, first.split('-'))
    while f'{y:04}-{m:02}' <= last:
        yield f'{y:04}-{m:02}'
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)

def read_gold(root):
    raw = root / 'research/gold-reserves/raw'
    data = json.loads((raw/'imf-dbnomics-volume.json').read_text())['series']['docs'][0]
    assert data['dimensions'] == {'FREQ':'M','INDICATOR':'RAFAGOLDV_OZT','REF_AREA':'PL','REF_SECTOR':'S1X'}
    volume = {p:(Decimal(str(v)), 'gold-imf-volume', {'sourceValue':v,'sourceUnit':'million fine troy ounces'}) for p,v in zip(data['period'],data['value']) if isinstance(v,(int,float))}
    # UNSD supplies the few early months absent from this IMF snapshot.
    soup = BeautifulSoup((raw/'un-gold-poland.html').read_text(encoding='utf8'), 'html.parser')
    assert 'millions of fine troy ounces end of period' in soup.get_text(' ',strip=True)
    for row in soup.find_all('tr'):
        cells=[c.get_text(' ',strip=True) for c in row.find_all(['td','th'],recursive=False)]
        cells=[cell for cell in cells if cell]
        if len(cells)<2 or not re.fullmatch(r'\d{4}(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)',cells[0]): continue
        p=cells[0][:4]+'-'+str(list(calendar.month_abbr).index(cells[0][4:].title())).zfill(2)
        if p not in volume and p<'2001-01':volume[p]=(Decimal(cells[1]),'gold-un-volume',{'sourceValue':cells[1],'sourceUnit':'million fine troy ounces'})
    evidence=[]
    for filename, sid in [('964irfcl-21-07-2026.xlsx','gold-nbp-june'),('gus-nbp-2026-08.xlsx','gold-nbp-july')]:
        ws=openpyxl.load_workbook(raw/filename,data_only=True)['Mon. Auth & Ctr. Gov']
        assert ws['E4'].value=='Poland, Rep. of' and ws['E6'].value=='US Dollars' and ws['E7'].value=='Millions'
        assert ws['C26'].value=='MCG_RAFAGOLDV_OZT'
        year,month=ws['E12'].value.split('M');p=f'{int(year):04}-{int(month):02}'
        v=Decimal(str(ws['E26'].value));volume[p]=(v,sid,{'sourceValue':float(v),'sourceUnit':'million fine troy ounces','sourceCell':'Mon. Auth & Ctr. Gov!E26'})
        evidence.append({'period':p,'sourceId':sid,'ouncesMillions':float(v),'tonnes':float(v*TONNES_PER_MILLION_OZT),'goldUsdMillions':ws['E25'].value,'reservesUsdMillions':ws['E13'].value})
    reports=reported_volume(root)
    report_sources={s['id']:s for s in reports['sources']}
    seen=set()
    for item in reports['observations']:
        p=item['period']; assert p not in seen, f'Duplicate reported month: {p}'
        seen.add(p)
        value=Decimal(item['value']);sid=item['sourceId'];assert value>0 and sid in report_sources
        # Direct source files take precedence; never silently overwrite them with press reports.
        if p in volume:assert abs(volume[p][0]-value)<=Decimal('0.01'), f'Conflicting source: {p}'
        else:volume[p]=(value,sid,{'sourceValue':float(value),'sourceValueText':item['value'],'sourceUnit':reports['unit'],'sourceDecimals':max(0,-value.as_tuple().exponent),'evidenceType':'secondary-report-of-NBP-observation','publishedAt':report_sources[sid]['publishedAt'],'sourceUrl':report_sources[sid]['url'],'corroborationUrl':item.get('corroborationUrl'),'interpolated':False})
    financial={}
    for row in csv.DictReader((raw/'ecb-reserves.csv').open()):
        assert row['REF_AREA']=='PL' and row['UNIT_MULT']=='6'
        if row['INSTR_ASSET']=='F11' and row['UNIT_MEASURE'] in ['PLN','EUR']:key='gold-'+row['UNIT_MEASURE'].lower()
        elif row['INSTR_ASSET']=='F' and row['CURRENCY_DENOM']=='X1' and row['UNIT_MEASURE']=='PLN':key='reserves-pln'
        else:continue
        if row['OBS_VALUE']!='':financial.setdefault(key,{})[row['TIME_PERIOD']]=(Decimal(row['OBS_VALUE']),row)
    return volume,financial,evidence

def add_gold_topic(root, topic, series, point, source, topics, sources):
    raw=root/'research/gold-reserves/raw'; volume,financial,evidence=read_gold(root)
    specs=[('gold-imf-volume','imf-dbnomics-volume.json','MFW / DBnomics · IRFCL · dokładna historia ilości','MFW / DBnomics'),('gold-un-volume','un-gold-poland.html','ONZ UNSD · Gold reserves · Poland','ONZ'),('gold-nbp-june','964irfcl-21-07-2026.xlsx','NBP / GUS SDDS · stan na 2026-06','NBP / GUS'),('gold-nbp-july','gus-nbp-2026-08.xlsx','NBP / GUS SDDS · stan na 2026-07','NBP / GUS'),('gold-ecb-values','ecb-reserves.csv','EBC · RAS · wycena złota i rezerw Polski','EBC')]
    for sid,filename,name,short in specs:
        receipt=json.loads((raw/(filename+'.receipt.json')).read_text())
        sources[sid]={**receipt,'id':sid,'name':name,'shortName':short,'license':'Original provider terms; source attribution required.','note':'Stan na koniec miesiąca; data publikacji/pobrania jest osobnym polem.'}
    reports=reported_volume(root)
    report_bytes=(raw/'reported-monthly-volume.json').read_bytes()
    for s in reports['sources']:
        sources[s['id']]={**s,'retrievedAt':reports['reviewedAt'],'reviewedAt':reports['reviewedAt'],'retrievalPrecision':'day','sourceType':'secondary-report','originalProvider':'NBP','extractionFile':'reported-monthly-volume.json','sha256':hashlib.sha256(report_bytes).hexdigest(),'bytes':len(report_bytes),'note':s['finding'],'license':'Extracted numerical facts with attribution; article text is not redistributed.'}
    topic(TOPIC,'Polska gromadzi złoto',COVERAGE,'partial')
    t=topics[TOPIC];t.update(titleEn='Poland builds its gold reserves',noteEn='Monthly gold quantities are complete from Jan 2000 through Aug 2026. Former gaps were filled with NBP observations reported by financial publications; source and precision are retained. No interpolation or estimates derived from gold prices. Financial valuations still have an ECB gap in Jan–Aug 2014. Includes NBP gold held abroad, not geological deposits.',methodologyPath='/stories/POLAND-GOLD.md',reelNote={'pl':'stan na koniec miesiąca','en':'month-end holdings'},defaultStart='2000-01-01',variants=['Ile ton złota zgromadziła Polska? Ciągły miesięczny wykres od 2000 r. do sierpnia 2026.','Ilość kontra wartość: tony i mld PLN, w oddzielnych skalach lub jako indeks 100. Wzrost wyceny nie oznacza zakupów.','Jaką część rezerw stanowi złoto? Udział procentowy według wartości na koniec miesiąca.','Zmiana zasobu miesiąc do miesiąca: słupki dodatnie i ujemne. Precyzja różnic zależy od dokładności publikowanych stanów.'])
    def make(key,pl,en,unit,unit_key,notes):
        s=series(TOPIC,key,'Polska · '+pl,unit,unit_key,'monthly',notes,'derived',entity='Poland',metric=en)
        s.update(countryCode='POL',metricLabels={'pl':pl,'en':en},includeEmptyPeriods=True)
        return s
    holdings=make('tonnes','Rezerwy złota','Gold reserves','tony','tonnes-gold','Miliony uncji czystego złota × 31,1034768 = tony metryczne. Dokładność ograniczona do precyzji źródła. '+COVERAGE)
    change=make('monthly-change','Zmiana rezerw złota m/m','Monthly change in gold reserves','tony / miesiąc','tonnes-gold-month','Różnica stanów dwóch sąsiednich miesięcy. Nie jest samodzielnym pomiarem zakupów netto; obejmuje również inne zmiany stanu i zaokrąglenia. '+COVERAGE)
    periods=list(months(min(volume),max(volume)))
    for i,p in enumerate(periods):
        entry=volume.get(p);sid=entry[1] if entry else 'gold-imf-volume'
        point(holdings,p,float(entry[0]*TONNES_PER_MILLION_OZT) if entry else None,sid,'month',**(entry[2] if entry else {'missingReason':'No exact monthly quantity in downloaded sources'}))
        previous=volume.get(periods[i-1]) if i else None
        delta=float((entry[0]-previous[0])*TONNES_PER_MILLION_OZT) if entry and previous else None
        point(change,p,delta,sid,'month',additionalSourceIds=[previous[1]] if previous else [],formula='tonnes(t) - tonnes(t-1)',missingReason=None if delta is not None else 'Needs both consecutive monthly observations')
    for key,pl,en,unit,uk in [('gold-pln','Wartość złota w PLN','Gold value in PLN','mld PLN','billion-pln'),('gold-eur','Wartość złota w EUR','Gold value in EUR','mld EUR','billion-eur'),('reserves-pln','Oficjalne aktywa rezerwowe','Official reserve assets','mld PLN','billion-pln')]:
        s=make(key,pl,en,unit,uk,'Wartość na koniec miesiąca, ceny bieżące. Miliony w źródle podzielono przez 1000. Zmiana wyceny zależy także od ceny złota i kursów walut; nie jest wolumenem zakupów. Brak danych EBC od 2014-01 do 2014-08.')
        for p in months(min(financial[key]),max(financial[key])):
            entry=financial[key].get(p)
            if entry:
                value,row=entry
                point(s,p,float(value/1000),'gold-ecb-values','month',seriesKey=row['KEY'],sourceValue=str(value),observationStatus=row['OBS_STATUS'])
            else:point(s,p,None,'gold-ecb-values','month',missingReason='No observation in ECB snapshot')
    share=make('share','Udział złota w rezerwach','Gold share of reserves','% rezerw','percent-reserves','100 × wartość złota w PLN / oficjalne aktywa rezerwowe w PLN, ten sam miesiąc i waluta. Zmiany wynikają również z wyceny oraz zmian pozostałych rezerw. Brak danych EBC od 2014-01 do 2014-08.')
    for p in months(min(financial['gold-pln']),max(financial['gold-pln'])):
        gold=financial['gold-pln'].get(p);total=financial['reserves-pln'].get(p)
        if gold and total and total[0]>0:point(share,p,float(100*gold[0]/total[0]),'gold-ecb-values','month',formula='100 * gold_PLN / total_reserves_PLN',goldValueMillions=float(gold[0]),totalValueMillions=float(total[0]))
        else:point(share,p,None,'gold-ecb-values','month',missingReason='Needs both same-month valuations')
    t['defaults']=[holdings['id']]
    last=volume[periods[-1]]
    audit={'volumeRange':[periods[0],periods[-1]],'missingVolumeMonths':[p for p in periods if p not in volume],'latestDirectNbpForm':evidence[-1],'latestObservation':{'period':periods[-1],'ouncesMillions':float(last[0]),'tonnes':float(last[0]*TONNES_PER_MILLION_OZT),'sourceId':last[1],'evidenceType':last[2].get('evidenceType','direct-source')},'nbpChecks':evidence,'reportedSupplementMonths':[p for p in periods if p in volume and volume[p][2].get('evidenceType')=='secondary-report-of-NBP-observation'],'interpolatedObservations':0,'rejectedEcbVolume':{'period':'2026-07','millionOunces':21,'reason':'Rounded to whole million ounces; source NBP template gives 20.583. Do not use ECB volume for precise tonnes or monthly changes.'}}
    audit['financialCoverage']={key:{'start':min(d),'end':max(d),'observations':len(d),'missingMonths':[p for p in months(min(d),max(d)) if p not in d]} for key,d in financial.items()}
    (root/'research/gold-reserves/quality.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8')
